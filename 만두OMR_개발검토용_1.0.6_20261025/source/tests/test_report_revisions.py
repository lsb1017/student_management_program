from copy import deepcopy
import pytest
from localomr.examdata import estimated_percentile
from localomr.report_math import report_history
from localomr.reports import report_pdf
from localomr.ranking import ranking_xlsx
from test_math import store,scan,exam,KEY,CONFIG


@pytest.mark.parametrize('score,expected',[(0,0),(250,2),(500,4),(1500,11),(2500,23),(3600,40),(5600,60),(6600,77),(7400,89),(8300,96),(7850,92.5),(9150,98),(10000,100)])
def test_cut_based_reference_anchors_and_interpolation(score,expected):
    assert estimated_percentile(score,10000,[83,74,66,56,36,25,15,5])==pytest.approx(expected)


def test_missing_invalid_and_endpoint_cuts():
    assert estimated_percentile(5000,10000,[]) is None
    assert estimated_percentile(5000,10000,[90,90,70,60,50,40,30,20]) is None
    assert estimated_percentile(10001,10000,[90,80,70,60,50,40,30,20]) is None
    cuts=[100,80,70,60,50,40,30,0]
    assert estimated_percentile(10000,10000,cuts)==96
    assert estimated_percentile(0,10000,cuts)==4


def test_new_results_do_not_change_estimate_and_cut_edits_do(store):
    s,student=store;eid=exam(s);answers=KEY.copy();answers[0]=None
    sid=scan(s,eid,student,answers=answers)
    original=s.results(eid)[0]['estimated_percentile']
    other=s.save_student('22222','추가학생',s.student(student)['class_id']);scan(s,eid,other,unique='other')
    assert next(r for r in s.results(eid) if r['id']==sid)['estimated_percentile']==original
    cfg=deepcopy(CONFIG);cfg['cuts']['1']=[90,80,70,60,50,40,30,20]
    s.save_exam('수학시험','시리즈','수학','1','2026-10-18',KEY,[3]*30,eid,cfg)
    assert s.scan(sid)['status']=='confirmed'
    assert next(r for r in s.results(eid) if r['id']==sid)['estimated_percentile']==pytest.approx(93.9)
    assert next(r for r in s.results(eid) if r['id']==sid)['grade']==2


def test_all_history_saved_latest_six_printed_and_no_counts_exported(store,tmp_path):
    import pypdfium2 as pdfium
    import openpyxl
    s,student=store
    for i in range(1,10):
        eid=s.save_exam(f'누적검증 {i}회','시리즈','수학',f'회차{i:02d}',f'2026-10-{i:02d}',KEY,[3]*30,config=CONFIG)
        scan(s,eid,student,unique=f'history{i}')
    records=s.results(student_id=student);assert len(records)==9
    assert [r['round'] for r in report_history(s,records[0])]==[f'회차{i:02d}' for i in range(4,10)]
    assert [r['round'] for r in report_history(s,records[4])]==[f'회차{i:02d}' for i in range(1,6)]
    pdf=tmp_path/'private.pdf';report_pdf(pdf,s,[records[0]['id']])
    with pdfium.PdfDocument(str(pdf)) as doc:
        text=''
        for page in doc:
            tp=page.get_textpage();text+=tp.get_text_range();tp.close();page.close()
    assert '추정 백분위' in text
    assert '공통 / 선택과목 점수' in text and '공통 66점 / 확률과 통계 24점' in text
    for removed in ('집단 최고 원점수','같은 선택과목 확정 성적 기준','오답 행 음영','영역별 100점 환산','반 평균','추정 백분위는 등급컷','공식 백분위가 아닙니다'):assert removed not in text
    for forbidden in ('집단 백분위','계산 인원','명)','회차01','회차02','회차03'):assert forbidden not in text
    for name in ('내 점수','평균 점수','순번','회차04','회차09'):assert name in text
    assert '최근 6회' not in text
    xlsx=tmp_path/'private.xlsx';ranking_xlsx(xlsx,s,eid,[s.student(student)['class_id']])
    book=openpyxl.load_workbook(xlsx);sheet=book.active
    assert sheet.max_column==17 and sheet['N8'].value=='추정 백분위' and sheet['N9'].value==100
    assert not any('인원' in str(c.value) for row in sheet for c in row)
    book.close()


def test_report_brand_setting_and_edited_area_names_follow_each_elective(store,tmp_path):
    import pypdfium2 as pdfium
    s,student=store;eid=exam(s);sid=scan(s,eid,student)
    other_class=s.classroom(s.academy('두번째학원'),'미적반');other=s.save_student('55551','미적학생',other_class);mid=scan(s,eid,other,'2','math2')
    assert s.report_name()=='양지용 수학연구실'
    s.set_report_name('설정한 연구실');backup=tmp_path/'brand.omrbackup';s.backup(backup);s.set_report_name('다른 연구실');s.restore(backup);assert s.report_name()=='설정한 연구실'
    cfg=deepcopy(CONFIG);cfg['topics']=['새지수']*10+['새삼각']*12+['새확률']*8;cfg['variants']['2']['topics']=['새미분']*8
    s.save_exam('수학시험','시리즈','수학','1','2026-10-18',KEY,[3]*30,eid,cfg)
    for scan_id,expected,absent in ((sid,'새확률','새미분'),(mid,'새미분','새확률')):
        path=tmp_path/f'{scan_id}.pdf';report_pdf(path,s,[scan_id])
        with pdfium.PdfDocument(str(path)) as doc:
            page=doc[0];tp=page.get_textpage();text=tp.get_text_range();tp.close();page.close()
        assert text.count('설정한 연구실')==1 and '수능과 가장 가까운 수학!' in text and '수학은 양지용입니다.' in text
        assert '정답 버전' not in text and '6등급 이하' not in text and '그 미만' not in text
        assert ('시험학원' if scan_id==sid else '두번째학원') in text
        assert text.count('새지수')>=2 and text.count('새삼각')>=2 and text.count(expected)>=2 and absent not in text
        positions=[text.index(header) for header in ['학생명','학생번호','학교','시험일','학원명','반']]
        assert positions==sorted(positions)
        for subtitle in ('학생정보','성적','단원별 성취도 분석'):assert subtitle in text
