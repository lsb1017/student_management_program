from datetime import datetime
from collections import Counter
from .examdata import ELECTIVES,percentile_label


def sorted_results(rows,field='name',descending=False,group_classes=False):
    if field not in ('name','score'):raise ValueError('정렬 기준을 확인하세요.')
    # Keep tied scores deterministic; the value of a rank never depends on print order.
    result=sorted(rows,key=lambda r:(r['snapshot']['name'],r['snapshot']['code'],r['id']))
    result.sort(key=(lambda r:r['snapshot']['name']) if field=='name' else (lambda r:r['score']),reverse=descending)
    if group_classes:result.sort(key=lambda r:(r['snapshot'].get('academy',''),r['snapshot'].get('classroom',''),r['snapshot'].get('class_id',0)))
    return result


def ranked_results(store,eid,class_ids):
    if not class_ids:raise ValueError('출력할 반을 하나 이상 체크하세요.')
    all_rows=store.results(eid);selected=[r for r in all_rows if r['snapshot'].get('class_id') in class_ids]
    def ranking(cohort):
        counts=Counter(r['score'] for r in cohort);below=0;ranks={}
        for score in sorted(counts,reverse=True):ranks[score]=below+1;below+=counts[score]
        return ranks
    overall=ranking(all_rows);filtered=ranking(selected)
    by_class={cid:[r for r in all_rows if r['snapshot'].get('class_id')==cid] for cid in class_ids}
    by_choice={choice:[r for r in all_rows if r['elective']==choice] for choice in {r['elective'] for r in selected}}
    for r in selected:
        classroom=by_class[r['snapshot']['class_id']];choice=by_choice[r['elective']]
        r['overall_rank']=overall[r['score']];r['selected_rank']=filtered[r['score']]
        r['class_rank']=ranking(classroom)[r['score']];r['elective_rank']=ranking(choice)[r['score']]
        r['overall_n']=len(all_rows);r['selected_n']=len(selected);r['class_n']=len(classroom);r['elective_n']=len(choice)
        # Percentile estimates are provided by Store.results, independent of cohorts.
    return selected


def ranking_xlsx(path,store,eid,class_ids,field='score',descending=True,group_classes=False):
    """App export: typed values are a snapshot, with literal text (no formulas)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins
    rows=sorted_results(ranked_results(store,eid,class_ids),field,descending,group_classes)
    if not rows:raise ValueError('선택한 반에 확정 성적이 없습니다.')
    exam=store.exam(eid);book=Workbook();sheet=book.active;sheet.title='학생별 석차'
    names=sorted({r['snapshot']['academy']+' / '+r['snapshot']['classroom'] for r in rows})
    notes=[exam['title']+' · 석차',f"시행일 {exam['day']} · {exam['subject']} · {exam['round']}회 · 원점수 기준",'선택 반: '+', '.join(names),
           '전체 석차는 시험 전체, 선택 반 석차는 체크한 반, 반 석차는 응시 당시 반, 과목 석차는 같은 선택과목 기준입니다.',
           '동점 석차는 1, 1, 3 방식입니다. 추정 백분위는 등급컷 기준 선형 환산 참고값이며 공식 백분위가 아닙니다.',
           f"출력 {datetime.now():%Y-%m-%d %H:%M} · 저장 시점의 값이며 이후 변경은 다시 내보내세요."]
    for i,note in enumerate(notes,1):
        sheet.cell(i,1,note).data_type='s';sheet.merge_cells(start_row=i,start_column=1,end_row=i,end_column=17)
        sheet.cell(i,1).font=Font(name='맑은 고딕',size=16 if i==1 else 10,bold=i==1,color='173C4B');sheet.row_dimensions[i].height=28 if i==1 else 22
    headers=['시험 전체 석차','선택 반 석차','반 석차','선택과목 석차','학생명','학생번호','학교','학원','반','선택과목','원점수','만점','등급','추정 백분위','공통 점수','선택 점수','오답 문항']
    for col,head in enumerate(headers,1):
        cell=sheet.cell(8,col,head);cell.font=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173C4B');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    sheet.row_dimensions[8].height=32
    for row_number,r in enumerate(rows,9):
        s=r['snapshot'];math=r['config'].get('form')=='math30'
        common=sum(p for i,(a,k,p) in enumerate(zip(r['answers'],r['key'],r['points'])) if i<22 and a==k)/100 if math else None
        choice=r['score']/100-common if math else None
        values=[r['overall_rank'],r['selected_rank'],r['class_rank'],r['elective_rank'],s['name'],s['code'][-5:] if math else s['code'],s.get('school',''),s['academy'],s['classroom'],ELECTIVES.get(r['elective'],r['subject']),r['score']/100,r['maximum']/100,r['grade'],r['estimated_percentile'] if r['estimated_percentile'] is not None else ('40 미만' if len(r['cuts'])==5 else None),common,choice,', '.join(map(str,r['wrong']))]
        for col,value in enumerate(values,1):
            cell=sheet.cell(row_number,col,value)
            if isinstance(value,str):cell.data_type='s'
            cell.font=Font(name='맑은 고딕',size=10,color='263549');cell.fill=PatternFill('solid',fgColor='F0F6F5' if row_number%2 else 'FFFFFF');cell.alignment=Alignment(vertical='center')
            if col in (11,12,15,16):cell.number_format='0.##'
            if col==14:cell.number_format='0.0'
            if col==6:cell.number_format='@'
        sheet.row_dimensions[row_number].height=23
    widths=[14,12,10,14,16,13,23,24,24,20,12,12,10,14,13,13,35]
    for i,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(i)].width=width
    sheet.freeze_panes='A9';sheet.auto_filter.ref=f'A8:Q{8+len(rows)}';sheet.sheet_view.showGridLines=False
    sheet.print_options.horizontalCentered=True;sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:8';sheet.print_area=f'A1:Q{8+len(rows)}';sheet.page_margins=PageMargins(left=.25,right=.25,top=.3,bottom=.3)
    book.save(path);book.close()
    return len(rows)
