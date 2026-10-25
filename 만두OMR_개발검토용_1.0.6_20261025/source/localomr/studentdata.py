"""Student workbook input and shared contact/code validation."""
import re
import sys
from pathlib import Path
from openpyxl import load_workbook

HEADERS=['성명','학생 연락처','학교','학부모 연락처','학원','반']
FIELDS=['name','phone','school','parent_phone','academy','classroom']

def phone_text(value,required=True):
    if value is None:value=''
    if isinstance(value,bool):raise ValueError('연락처를 숫자로 입력하세요.')
    if isinstance(value,(int,float)):
        if value!=int(value):raise ValueError('연락처에 소수를 사용할 수 없습니다.')
        value=str(int(value))
        if required and len(value)<=4:value=value.zfill(4)
    value=str(value).strip()
    if not value and not required:return ''
    if not re.fullmatch(r'[0-9() +\-]+',value):raise ValueError('연락처는 숫자와 하이픈으로 입력하세요.')
    digits=re.sub(r'[^0-9]','',value)
    if required and len(digits)==4:return digits
    if 9<=len(digits)<=11 and digits.startswith('0'):return digits
    raise ValueError('학생 연락처는 뒤 4자리 또는 전체 번호를 입력하세요. 전체 번호의 첫 0을 확인하세요.' if required else '학부모 연락처는 첫 0을 포함한 전체 번호로 입력하세요.')

def validate_student(row):
    r=dict(row);r['name']=str(r.get('name') or '').strip()
    if not r['name'] or len(r['name'])>60:raise ValueError('성명은 필수이며 60자 이내로 입력하세요.')
    r['phone']=phone_text(r.get('phone'));r['parent_phone']=phone_text(r.get('parent_phone'),False)
    for field in ('school','academy','classroom'):
        r[field]=str(r.get(field) or '').strip()
        if len(r[field])>(800 if field=='classroom' else 80):raise ValueError('학교·학원·반은 80자 이내로 입력하세요.')
    r['classrooms']=[x.strip() for x in r['classroom'].split(';') if x.strip()]
    if any(len(x)>80 for x in r['classrooms']):raise ValueError('반 이름은 80자 이내로 입력하세요.')
    return r

def next_code(tail,used):
    for suffix in range(1,10):
        code='0'+tail+str(suffix)
        if code not in used:return code
    raise ValueError(f'{tail}: 구분번호 1~9를 모두 사용 중입니다. 다른 연락처 뒤 4자리를 입력하세요.')

def display_code(code):
    return code[1:] if len(code)==6 and code.startswith('0') else code

def plan_students(store,rows):
    existing=store.students();used={s['code'] for s in existing};plan=[]
    for i,row in enumerate(rows,1):
        r=dict(row);line=r.get('row',i)
        try:
            r=validate_student(r)
            def same(s):
                phone=s.get('phone') or (s['code'][1:5] if s.get('code','').startswith('0') else '')
                return s['name']==r['name'] and phone[-4:]==r['phone'][-4:] and (phone==r['phone'] or len(phone)==4 or len(r['phone'])==4)
            matches=[s for s in existing if same(s)]
            target=str(r.get('student_code') or '').strip()
            if target:
                target=target.zfill(6)
                targets=[x for x in existing if x.get('code')==target]
                if len(targets)!=1 or targets[0]['name']!=r['name']:raise ValueError('기존 학생번호와 성명이 일치하지 않습니다.')
                r.update(action='link',code=target,student_id=targets[0]['id'],message=f'{line}행: 기존 학생에 반 추가 (성적·기존 정보 유지)')
            elif matches:
                r.update(action='skip',code=matches[0]['code'],message=f'{line}행: 같은 성명·연락처가 있어 건너뜀 (기존 정보 유지)')
            else:
                r.update(action='add',code=next_code(r['phone'][-4:],used),message=f'{line}행: 신규 등록')
                used.add(r['code']);existing.append(r)
        except ValueError as exc:r.update(action='error',code='',message=f'{line}행: {exc}')
        plan.append(r)
    return plan

def read_students(path):
    path=Path(path)
    if path.suffix.lower()!='.xlsx':raise ValueError('학생 등록 서식(.xlsx)을 선택하세요.')
    if path.stat().st_size>20*1024*1024:raise ValueError('20MB 이하의 엑셀을 선택하세요.')
    book=load_workbook(path,read_only=True,data_only=False)
    try:
        sheet=book['학생등록'] if '학생등록' in book.sheetnames else book.worksheets[0]
        if sheet.max_row>10020:raise ValueError('한 번에 10,000명까지 등록할 수 있습니다.')
        header=None;rows=[]
        for index,cells in enumerate(sheet.iter_rows(max_col=7),1):
            vals=[c.value for c in cells]
            if header is None:
                names=[str(v or '').replace(' ','').replace('*','').replace('(필수)','') for v in vals]
                if names[:6]==[h.replace(' ','') for h in HEADERS]:header=index
                elif index>=20:raise ValueError('성명 / 학생 연락처 / 학교 / 학부모 연락처 / 학원 / 반 순서의 서식을 사용하세요.')
                continue
            if not any(v is not None and str(v).strip() for v in vals):continue
            if any(c.data_type=='f' for c in cells):raise ValueError(f'{index}행: 수식 대신 값을 붙여넣어 주세요.')
            rows.append(dict(zip(FIELDS+['student_code'],vals),row=index))
        if not header or not rows:raise ValueError('등록할 학생이 없습니다. 서식의 입력란을 채워 주세요.')
        return rows
    finally:book.close()

def template_path():
    root=Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parent.parent))
    return root/'assets'/'학생_일괄등록_서식.xlsx'
