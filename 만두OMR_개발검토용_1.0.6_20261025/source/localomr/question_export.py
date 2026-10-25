"""Use the same canonical columns as the question workbook importer."""
from pathlib import Path
import os,uuid
from .examdata import effective,ELECTIVES
from .question_import import HEADERS

def export_question_book(path,exam):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    workbook=Workbook();workbook.remove(workbook.active)
    math=exam['config'].get('form')=='math30'
    choices=[k for k in ELECTIVES if effective(exam,k)['configured']] if math else [None]
    for choice in choices:
        data=effective(exam,choice);sheet=workbook.create_sheet(ELECTIVES[choice] if math else '정답표')
        sheet.append(HEADERS)
        for i,key in enumerate(data['key']):
            sheet.append([i+1,data['topics'][i],data['areas'][i],data['skills'][i],data['difficulties'][i],data['points'][i]/100,key,data['kinds'][i],data['conditions'][i]])
        for row in sheet:
            for cell in row:
                # Text that resembles an Excel formula is still literal source data.
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='center',horizontal='left' if cell.column in (2,3,4) else 'center')
                cell.font=Font(name='맑은 고딕',size=10,bold=cell.row==1,color='49395D')
                if cell.row==1:cell.fill=PatternFill('solid',fgColor='EEE7F8')
        for col,width in enumerate([9,23,19,42,12,10,10,15,13],1):sheet.column_dimensions[get_column_letter(col)].width=width
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    if not choices:raise ValueError('다운로드할 정답표가 없습니다.')
    path=Path(path);temp=path.with_name(path.name+'.'+uuid.uuid4().hex+'.tmp')
    try:workbook.save(temp);os.replace(temp,path)
    finally:workbook.close();temp.unlink(missing_ok=True)
    return path
