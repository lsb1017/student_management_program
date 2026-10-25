"""Read-only Excel question import, validated before it reaches an exam draft."""
from dataclasses import dataclass,field
from pathlib import Path
from decimal import Decimal,InvalidOperation
import re,zipfile
from .examdata import NUMERIC,ELECTIVES
from .store import points_cents

HEADERS=['문항','영역분류','영역','유형','난이도','배점','정답','문제형태','채점조건']

def header_name(value):
    value=re.sub(r'\s+','',clean(value))
    return {'단원':'영역분류','과목':'영역','학습내용':'유형'}.get(value,value)


@dataclass
class QuestionSheet:
    name:str
    rows:list=field(default_factory=list)
    errors:list=field(default_factory=list)
    elective:str|None=None
    group:str=''

    @property
    def maximum(self):return sum(points_cents(r['points']) for r in self.rows)/100


def clean(value):return str(value).strip() if value is not None else ''


def integer(value):
    if value is None or isinstance(value,bool):raise ValueError('정수가 필요합니다.')
    number=Decimal(str(value).strip())
    if not number.is_finite() or number!=number.to_integral_value():raise ValueError('정수가 필요합니다.')
    return int(number)


def parse_rows(name,values,required_questions=None,require_elective=True):
    out=QuestionSheet(name);header=None;indices={};seen=set()
    for offset,row in enumerate(values[:20]):
        normalized=[header_name(v) for v in row]
        if all(h in normalized for h in HEADERS):
            header=offset;indices={h:normalized.index(h) for h in HEADERS};break
    if header is None:out.errors.append('문항·단원·과목·학습내용·난이도·배점·정답·문제형태·채점조건 머리글을 찾지 못했습니다. 기존 머리글도 지원합니다.');return out
    for excel_row,values_row in enumerate(values[header+1:],header+2):
        if not any(clean(v) for v in values_row):continue
        row={h:values_row[i] if i<len(values_row) else None for h,i in indices.items()}
        if any(clean(v).startswith('=') for v in row.values()):out.errors.append(f'{excel_row}행: 수식은 값으로 붙여넣은 뒤 불러오세요.');continue
        try:
            q=integer(row['문항']);answer=integer(row['정답']);points=points_cents(row['배점'])/100
            if q in seen:raise ValueError(f'{q}번이 중복되었습니다.')
            if not 1<=q<=30:raise ValueError('수학 OMR 문항 번호는 1~30입니다.')
            seen.add(q);numeric=q in NUMERIC
            if not (0<=answer<=999 if numeric else 1<=answer<=5):raise ValueError('객관식은 1~5, 숫자형은 0~999 정답이 필요합니다.')
            kind=clean(row['문제형태']).replace(' ','')
            accepted=('숫자답안','숫자단답','숫자단답형','주관식','숫자형') if numeric else ('선다형','객관식','5지선다','객관식5지선다')
            if kind not in accepted:raise ValueError(f'{q}번의 문제형태가 OMR 답란과 맞지 않습니다.')
            condition=clean(row['채점조건']).lower() or 'and'
            if condition!='and':raise ValueError('현재는 단일 정답(and) 채점조건을 지원합니다.')
            out.rows.append({'q':q,'key':answer,'points':points,'topic':clean(row['영역분류']) or '미분류','area':clean(row['영역']),'skill':clean(row['유형']),'difficulty':clean(row['난이도']),'condition':condition,'kind':kind})
        except (ValueError,InvalidOperation) as exc:out.errors.append(f'{excel_row}행: {exc}')
    out.rows.sort(key=lambda r:r['q'])
    required=set(range(1,31)) if required_questions is None else set(required_questions)
    missing=sorted(required-set(r['q'] for r in out.rows))
    if missing:out.errors.append('누락/입력 오류 문항: '+', '.join(map(str,missing)))
    extra=sorted(set(r['q'] for r in out.rows)-required)
    if extra:out.errors.append('선택 구역 밖 문항: '+', '.join(map(str,extra)))
    if out.rows and out.maximum<=0:out.errors.append('총 배점이 0입니다.')
    names={re.sub(r'\s+','',r['area']) for r in out.rows if r['q']>=23}
    candidates=[k for k,v in ELECTIVES.items() if re.sub(r'\s+','',v) in names]
    if len(candidates)==1:out.elective=candidates[0]
    else:
        guesses=[k for k,word in [('1','확통'),('1','확률과통계'),('2','미적분'),('3','기하')] if word in name.replace(' ','')]
        if len(set(guesses))==1:out.elective=guesses[0]
    if not out.elective and require_elective:out.errors.append('선택과목을 판별할 수 없습니다. 시트명 또는 선택 문항의 영역을 확인하세요.')
    out.group=re.sub(r'(확률\s*과\s*통계|확통|미적분|기하)','',name).strip()
    return out


def read_question_book(path):
    import openpyxl
    path=Path(path)
    if path.suffix.lower()!='.xlsx' or path.stat().st_size>20*1024*1024:raise ValueError('20MB 이하의 .xlsx 파일을 선택하세요.')
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist())>100*1024*1024:raise ValueError('압축 해제 크기가 큰 엑셀 파일입니다. 필요한 시트만 나누어 저장하세요.')
    book=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        if len(book.sheetnames)>100:raise ValueError('시트 100개 이하로 나누어 주세요.')
        sheets=[]
        for sheet in book:
            if sheet.max_row>2000 or sheet.max_column>100:sheets.append(QuestionSheet(sheet.title,errors=['시트가 너무 큽니다. 문항 정보만 남긴 사본을 선택하세요.']));continue
            values=[list(r) for r in sheet.iter_rows(values_only=True)];sheets.append(parse_rows(sheet.title,values))
        return sheets
    finally:book.close()


def combine_sheets(sheets):
    if not sheets:raise ValueError('가져올 시트를 선택하세요.')
    if any(s.errors for s in sheets):raise ValueError('선택한 시트의 오류를 먼저 수정하세요.')
    if len({s.elective for s in sheets})!=len(sheets):raise ValueError('같은 선택과목의 시트를 두 개 이상 선택할 수 없습니다.')
    if len({s.group for s in sheets})!=1:raise ValueError('서로 다른 시험·회차의 시트가 선택되었습니다. 같은 시험 시트만 선택하세요.')
    common=sheets[0].rows[:22]
    fields=('key','points','topic','area','skill','difficulty','condition')
    for sheet in sheets[1:]:
        mismatches=[i+1 for i in range(22) if any(common[i][f]!=sheet.rows[i][f] for f in fields)]
        if mismatches:raise ValueError('선택과목 간 공통 문항 정보가 다릅니다: '+', '.join(map(str,mismatches)))
    return {s.elective:s.rows for s in sheets}
