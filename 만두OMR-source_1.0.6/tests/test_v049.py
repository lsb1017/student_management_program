import sqlite3
import threading
from pathlib import Path
import pytest
from PySide6.QtCore import QTimer, Qt
from PySide6.QtGui import QCloseEvent
from PySide6.QtWidgets import QFileDialog, QMessageBox, QProgressBar
from localomr.busy import run_store_task, BusyDialog, BUSY_TEXT
from localomr.store import Store
from localomr.ui import MainWindow
from localomr.reports import report_pdf
from localomr.ranking import ranking_xlsx
from test_math import store, exam, scan
from test_v043 import app


def test_backup_keeps_ui_alive_and_cannot_close_mid_write(store, tmp_path, app, monkeypatch):
    db,_=store;w=MainWindow(db);w.score_exam.setCurrentIndex(1);db.academy('변경');target=tmp_path/'saved.omrbackup'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,**k:(str(target),''))
    original=Store.backup;release=threading.Event();observed=[];worker_ids=[]
    def controlled(worker,path):
        worker_ids.append(threading.get_ident())
        assert release.wait(5), 'UI did not stay responsive'
        return original(worker,path)
    monkeypatch.setattr(Store,'backup',controlled)
    def inspect():
        dialog=app.activeModalWidget()
        if not isinstance(dialog,BusyDialog):return
        observed.append((dialog.message.text(),dialog.windowModality()))
        bar=dialog.findChild(QProgressBar)
        observed.append((bar.minimum(),bar.maximum()))
        dialog.reject();dialog.close()
        observed.append(dialog.isVisible())
        event=QCloseEvent();w.closeEvent(event);observed.append(event.isAccepted())
        release.set();timer.stop()
    timer=QTimer();timer.timeout.connect(inspect);timer.start(10)
    try:assert w.backup(other_location=True) is True
    finally:timer.stop();release.set()
    assert observed==[(BUSY_TEXT,Qt.WindowModality.ApplicationModal),(0,0),True,False]
    assert worker_ids and worker_ids[0]!=threading.get_ident()
    assert target.is_file() and not w.handoff.needs_backup() and not w._busy_operation
    with sqlite3.connect(target) as copied:
        assert copied.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert copied.execute("SELECT COUNT(*) FROM academies WHERE name='변경'").fetchone()[0]==1
    w.hide();w.deleteLater()


def test_worker_restore_preserves_safety_backup_and_rejects_bad_file(store,tmp_path,app):
    db,_=store;w=MainWindow(db)
    saved=run_store_task(w,db,lambda worker:worker.backup(tmp_path/'old.omrbackup'),readonly=True)
    db.set_report_name('현재 자료')
    safety=run_store_task(w,db,lambda worker:worker.restore(saved))
    assert db.report_name()!='현재 자료'
    with sqlite3.connect(safety) as before:
        assert before.execute("SELECT value FROM settings WHERE name='report_name'").fetchone()[0]=='현재 자료'
    bad=tmp_path/'bad.omrbackup';bad.write_bytes(b'not sqlite')
    before=db.report_name()
    with pytest.raises(sqlite3.DatabaseError):run_store_task(w,db,lambda worker:worker.restore(bad))
    assert db.report_name()==before and not w._busy_operation
    assert run_store_task(w,db,lambda worker:worker.report_name(),readonly=True)==before
    w.hide();w.deleteLater()


def test_worker_report_and_rank_exports_keep_data_and_order(store,tmp_path,app):
    db,student=store;eid=exam(db);sid=scan(db,eid,student);w=MainWindow(db)
    before=db.results();audit=db.db.execute('SELECT MAX(id) FROM audit').fetchone()[0]
    cid=db.student(student)['class_id']
    pdf=tmp_path/'report.pdf';xlsx=tmp_path/'rank.xlsx'
    run_store_task(w,db,lambda worker:report_pdf(pdf,worker,[sid]),readonly=True)
    assert run_store_task(w,db,lambda worker:ranking_xlsx(xlsx,worker,eid,[cid]),readonly=True)==1
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(str(pdf)) as doc:assert len(doc)==2
    from openpyxl import load_workbook
    book=load_workbook(xlsx,read_only=True);assert book.active.max_row>=2;book.close()
    assert db.results()==before and db.db.execute('SELECT MAX(id) FROM audit').fetchone()[0]==audit
    w.hide();w.deleteLater()


def test_worker_blocks_overlap_and_does_not_create_or_migrate_database(store,tmp_path,app):
    db,_=store;w=MainWindow(db);w.score_exam.setCurrentIndex(1);w.worker=object()
    with pytest.raises(ValueError):run_store_task(w,db,lambda worker:None)
    w.worker=None
    with pytest.raises(sqlite3.OperationalError):Store.open_existing(tmp_path/'missing.sqlite3')
    assert not (tmp_path/'missing.sqlite3').exists()
    before=db.db.execute('SELECT MAX(id) FROM audit').fetchone()[0]
    with pytest.raises(sqlite3.OperationalError):
        run_store_task(w,db,lambda worker:worker.academy('금지'),readonly=True)
    assert not w._busy_operation
    assert db.db.execute('SELECT MAX(id) FROM audit').fetchone()[0]==before
    w.hide();w.deleteLater()
