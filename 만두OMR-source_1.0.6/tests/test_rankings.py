import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from copy import deepcopy
import pytest
from localomr.store import Store
from localomr.omr import ReadResult
from localomr.examui import MATH_POINTS
from localomr.ranking import ranked_results,sorted_results,ranking_xlsx


@pytest.fixture
def setup(tmp_path):
    db=Store(tmp_path/'rank.sqlite3');academy=db.academy('학원');a=db.classroom(academy,'A반');b=db.classroom(academy,'B반')
    cuts={'1':[83,74,66,56,36,25,15,5],'2':[90,80,70,60,50,40,30,20]}
    config={'form':'math30','primary':'1','cuts':cuts,'variants':{'2':{'key':[1]*8,'points':MATH_POINTS[22:],'topics':['선택']*8,'skills':['']*8}}}
    eid=db.save_exam('석차 검증','시리즈','수학','1','2026-10-18',[1]*30,MATH_POINTS,config=config)
    for i,(name,cid,choice,wrong) in enumerate([('가학생',a,'1',[10,14,21,29]),('나학생',a,'2',[10,14,21,29]),('다학생',b,'1',[]),('라학생',b,'1',[8,10,14,21,29])]):
        student=db.save_student(str(i+1),name,cid);answers=[None if q in wrong else 1 for q in range(30)]
        source=db.db.execute('INSERT INTO sources(hash,name,data) VALUES(?,?,?)',(str(i),name,b'example')).lastrowid;db.db.commit()
        result=ReadResult(None,[{'answer':v,'flag':'clear','levels':[]} for v in answers],None);result.elective=choice
        sid,_=db.add_scan(eid,source,1,str(i),b'example',result);db.save_review(sid,student,answers,[True]*30,True,elective=choice);db.confirm(sid)
    yield db,eid,a,b
    db.close()


def test_ranks_keep_whole_exam_cohorts_and_subject_cuts(setup):
    db,eid,a,b=setup;rows=ranked_results(db,eid,[a]);assert len(rows)==2
    assert {r['overall_rank'] for r in rows}=={2}
    assert {r['selected_rank'] for r in rows}=={1}
    assert {r['class_rank'] for r in rows}=={1}
    assert {r['overall_n'] for r in rows}=={4}
    by_choice={r['elective']:r for r in rows}
    assert by_choice['1']['grade']==1 and by_choice['2']['grade']==2
    assert by_choice['1']['elective_rank']==2 and by_choice['2']['elective_rank']==1
    assert [r['overall_rank'] for r in sorted_results(ranked_results(db,eid,[a,b]),'score',True)]==[1,2,2,4]


@pytest.mark.parametrize('field,descending,expected',[
    ('name',False,['가학생','나학생','다학생','라학생']),('name',True,['라학생','다학생','나학생','가학생']),
    ('score',False,['라학생','가학생','나학생','다학생']),('score',True,['다학생','가학생','나학생','라학생'])])
def test_all_four_print_orders(setup,field,descending,expected):
    db,eid,a,b=setup;rows=ranked_results(db,eid,[a,b]);sorted_rows=sorted_results(rows,field,descending)
    assert [r['snapshot']['name'] for r in sorted_rows]==expected
    grouped=sorted_results(rows,'score',True,True)
    assert [r['snapshot']['class_id'] for r in grouped]==[a,a,b,b]


def test_xlsx_typed_values_no_formula_injection(setup,tmp_path):
    import openpyxl
    db,eid,a,b=setup
    row=db.results(eid)[0];snapshot=deepcopy(row['snapshot']);snapshot['name']='=HYPERLINK("invalid")'
    import json
    db.db.execute('UPDATE scans SET snapshot=? WHERE id=?',(json.dumps(snapshot),row['id']));db.db.commit()
    output=tmp_path/'ranks.xlsx';assert ranking_xlsx(output,db,eid,[a,b])==4
    book=openpyxl.load_workbook(output);sheet=book.active
    assert sheet.max_row==12 and sheet.freeze_panes=='A9'
    assert sheet['K9'].value==100 and sheet['K9'].data_type=='n'
    assert any(sheet.cell(i,5).value.startswith('=') and sheet.cell(i,5).data_type=='s' for i in range(9,13))
    assert all(sheet.cell(i,6).data_type=='s' for i in range(9,13))
    assert sheet.auto_filter.ref=='A8:Q12';book.close()
    with pytest.raises(ValueError):ranking_xlsx(tmp_path/'empty.xlsx',db,eid,[])


def test_selected_class_pdf_page_order(setup,tmp_path):
    import pypdfium2 as pdfium
    from localomr.reports import report_pdf
    db,eid,a,b=setup;rows=sorted_results(ranked_results(db,eid,[a]),'name',True)
    output=tmp_path/'batch.pdf';report_pdf(output,db,[r['id'] for r in rows])
    with pdfium.PdfDocument(str(output)) as doc:
        assert len(doc)==4
        for page_no,name in [(0,'나학생'),(2,'가학생')]:
            p=doc[page_no];t=p.get_textpage();assert name in t.get_text_range();t.close();p.close()


def test_batch_dialog_selection_and_order(setup):
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import Qt
    from localomr.batchui import BatchOutputDialog
    app=QApplication.instance() or QApplication([]);db,eid,a,b=setup;dialog=BatchOutputDialog(db,eid)
    assert len(dialog.rows)==4
    dialog.group.setChecked(False);dialog.field.setCurrentIndex(1);dialog.direction.setCurrentIndex(1)
    assert dialog.rows[0]['score']==10000
    dialog.classes.item(dialog.class_ids.index(b),0).setCheckState(Qt.CheckState.Unchecked)
    assert len(dialog.rows)==2 and all(r['snapshot']['class_id']==a for r in dialog.rows)
    dialog.check_all(False);assert not dialog.rows and not dialog.pdf_button.isEnabled();dialog.close()
