import json,sqlite3,hashlib
from copy import deepcopy
from pathlib import Path
import pytest
from PySide6.QtWidgets import QApplication,QMessageBox,QFileDialog
from PySide6.QtCore import Qt,QEvent
from PySide6.QtGui import QCloseEvent
from localomr.store import Store
from localomr.ui import MainWindow,STYLE
from localomr.backup_tools import counts,inspect_backup,comparison
from localomr.handoff import revision

@pytest.fixture
def fixture(tmp_path,monkeypatch):
    app=QApplication.instance() or QApplication([])
    app.sendPostedEvents(None,QEvent.Type.DeferredDelete)
    if app.style().objectName().lower()!='fusion':app.setStyle('Fusion')
    if app.styleSheet()!=STYLE:app.setStyleSheet(STYLE)
    db=Store(tmp_path/'data.sqlite3');a=db.academy('가상학원');c=db.classroom(a,'가상반')
    student=db.save_student_contact('가상학생','01012345678',class_id=c)
    retired=db.save_student_contact('가상퇴원생','9999',class_id=c,active=False)
    e=db.save_exam('가상 7회차','가상','수학','7','2026-10-20',[1,2,3],[2,3,4])
    warnings=[];monkeypatch.setattr(QMessageBox,'warning',lambda *args:warnings.append(str(args[2])))
    monkeypatch.setattr(QMessageBox,'information',lambda *args:QMessageBox.StandardButton.Ok)
    w=MainWindow(db)
    yield app,db,w,e,student,c,warnings
    w.hide();w.deleteLater();app.sendPostedEvents(None,QEvent.Type.DeferredDelete);app.processEvents();db.close()

def add_scan(db,e,student,tmp_path):
    from localomr.omr import ReadResult
    p=tmp_path/'image.png';p.write_bytes(b'synthetic-source')
    reading=ReadResult(None,[{'answer':a,'flag':'clear','levels':[]} for a in [1,2,3]],None)
    sid,_=db.add_scan(e,db.source(p),1,'synthetic',b'',reading)
    db.save_review(sid,student,[1,2,3],[False]*3,True)
    return sid

def test_startup_requires_explicit_choice_and_skip_is_session_only(fixture,monkeypatch):
    app,db,w,e,student,c,warnings=fixture
    assert w.score_exam.currentData() is None and w.scan_table.rowCount()==0
    seen=[]
    monkeypatch.setattr('localomr.popups.ask',lambda *args:seen.append(args) or QMessageBox.StandardButton.Cancel)
    w.nav.setCurrentRow(3);assert w.pages.currentIndex()==0 and len(seen)==1
    monkeypatch.setattr('localomr.popups.ask',lambda *args:QMessageBox.StandardButton.No)
    baseline=w.handoff.state['baseline'];w.nav.setCurrentRow(3)
    assert w.pages.currentIndex()==3 and w.startup_backup_ready and w.handoff.state['baseline']==baseline
    monkeypatch.setattr('localomr.popups.ask',lambda *args:pytest.fail('Must not repeat this run'))
    w.nav.setCurrentRow(0);w.nav.setCurrentRow(5)
    w.score_exam.setCurrentIndex(w.score_exam.findData(e));w.refresh();assert w.score_exam.currentData()==e
    assert w.score_exam.currentText()=='가상 7회차'
    db.archive_exams([e]);w.refresh();assert w.score_exam.currentData() is None
    second=MainWindow(db);assert not second.startup_backup_ready;second.hide();second.deleteLater()

def test_import_confirmation_after_file_selection_and_cancel_writes_nothing(fixture,monkeypatch):
    app,db,w,e,student,c,warnings=fixture
    w.score_exam.setCurrentIndex(w.score_exam.findData(e));w.score_academy.setCurrentIndex(1);w.score_class.setCurrentIndex(w.score_class.findData(c))
    order=[];monkeypatch.setattr(QFileDialog,'getOpenFileNames',lambda *a:order.append('file') or (['chosen.pdf'],''))
    monkeypatch.setattr('localomr.popups.ask',lambda *a:order.append(a[2]) or QMessageBox.StandardButton.Cancel)
    monkeypatch.setattr(w,'start_import',lambda *a:order.append('start'))
    before=revision(db);w.import_dialog();assert order==['file','채점을 시작할까요?'] and revision(db)==before
    order.clear();monkeypatch.setattr('localomr.popups.ask',lambda *a:order.append(a[2]) or QMessageBox.StandardButton.Ok)
    w.import_dialog();assert order==['file','채점을 시작할까요?','start']
    order.clear();monkeypatch.setattr(QFileDialog,'getOpenFileNames',lambda *a:([],''));w.import_dialog();assert not order

def test_backup_time_counts_restore_and_old_format(fixture,tmp_path,monkeypatch):
    app,db,w,e,student,c,warnings=fixture
    sid=add_scan(db,e,student,tmp_path);db.confirm(sid)
    path=tmp_path/'new.omrbackup';before=revision(db);db.backup(path)
    assert before==revision(db)
    incoming=inspect_backup(path);assert incoming['students']==2 and incoming['results']==1 and incoming['saved_at']
    assert '+1명' in comparison({'students':1,'results':3},incoming) and '-2건' in comparison({'students':1,'results':3},incoming)
    old=tmp_path/'old.omrbackup';src=sqlite3.connect(path);dst=sqlite3.connect(old);src.backup(dst);dst.execute('DROP TABLE backup_metadata');dst.commit();src.close();dst.close()
    assert inspect_backup(old)['saved_at'] is None
    db.save_student_contact('추가학생','7777',class_id=c)
    safety=db.restore(old);assert safety.exists() and counts(db.db)=={'students':2,'results':1}
    assert db.scan(sid)['status']=='confirmed' and db.scan(sid)['answers']==[1,2,3]
    assert not warnings

def test_readonly_archive_preserves_both_databases(fixture,tmp_path):
    from localomr.archiveui import prepare_archive,ArchiveDialog
    app,db,w,e,student,c,warnings=fixture
    sid=add_scan(db,e,student,tmp_path);db.confirm(sid);source=db.backup(tmp_path/'archive.omrbackup')
    digest=hashlib.sha256(source.read_bytes()).hexdigest();before=revision(db)
    path=tmp_path/'view.sqlite3';prepare_archive(source,path);archive=Store.open_existing(path,readonly=True)
    try:
        with pytest.raises(sqlite3.OperationalError):archive.academy('금지')
        dialog=ArchiveDialog(archive,source,w);assert dialog.grid.rowCount()==1;dialog.close()
    finally:archive.close()
    assert revision(db)==before and hashlib.sha256(source.read_bytes()).hexdigest()==digest

def test_unsaved_manual_guard_and_retired_default(fixture,monkeypatch):
    app,db,w,e,student,c,warnings=fixture
    assert w.manual.roster.rowCount()==1
    w.manual.include_inactive.setChecked(True);assert w.manual.roster.rowCount()==2
    w.manual.open_student(e,student);w.manual.full_marks()
    monkeypatch.setattr(QMessageBox,'exec',lambda box:0)
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a:pytest.fail('Unsaved drafts must be handled first'))
    assert w.backup() is False and w.manual.has_unsaved() and not db.results()

def test_answer_undo_and_stale_protection(fixture,tmp_path):
    app,db,w,e,student,c,warnings=fixture
    sid=add_scan(db,e,student,tmp_path);db.confirm(sid)
    w.score_exam.setCurrentIndex(w.score_exam.findData(e));pane=w.review_pane
    pane.update_answer(0,4);assert db.scan(sid)['status']=='pending' and db.scan(sid)['score']==700
    pane.undo_answer();assert db.scan(sid)['answers']==[1,2,3] and db.scan(sid)['score']==900 and db.scan(sid)['status']=='pending'
    pane.update_answer(1,5);db.confirm(sid);pane.undo_answer()
    assert db.scan(sid)['answers'][1]==5 and db.scan(sid)['status']=='confirmed'
    assert all('raw' not in item[2] and 'normalized' not in item[2] for item in pane.undo_history)

def test_status_filters_checkbox_selection_and_alignment(fixture,tmp_path):
    app,db,w,e,student,c,warnings=fixture
    sid=add_scan(db,e,student,tmp_path);w.score_exam.setCurrentIndex(w.score_exam.findData(e))
    assert '검수 대기 1장' in w.pending_mascot.text() and '미인식 0장' in w.scan_summary.text()
    assert w.scan_table.item(0,5).text()=='1'
    w.review_pane.update_answer(0,4);assert w.scan_table.item(0,5).text()=='4\n×'
    w.filter_scan_status('failed');assert w.scan_table.rowCount()==0
    w.filter_scan_status('pending');assert w.scan_table.rowCount()==1
    db.confirm(sid);w.refresh();w.result_table.set_checks(True)
    assert w.selection_count.text()=='선택 1건' and w.result_table.checked_ids()==[sid]
    w.result_search.setText('없는학생');assert w.selection_count.text()=='선택 0건'
    assert w.manual.roster.item(0,1).textAlignment() & Qt.AlignmentFlag.AlignHCenter

def test_storage_empty_state_and_keyboard_controls(fixture):
    from localomr.storageui import StorageDialog
    app,db,w,e,student,c,warnings=fixture
    dialog=StorageDialog(db,w);assert dialog.days.value()==7 and dialog.days.minimum()==0
    assert dialog.empty.text() and not dialog.clean.isEnabled()
    dialog.days.setValue(1);dialog.days.stepUp();assert dialog.days.value()==2
    dialog.broaden();assert dialog.days.value()==2 and dialog.exam.currentData() is None
    assert '정리 대상이 없어요' in dialog.empty.text();dialog.close()

def test_settings_keeps_group_drafts_and_progress_stops(fixture):
    from localomr.busy import ManduProgressBar
    app,db,w,e,student,c,warnings=fixture
    panel=w.cheer_panel;panel.editors['perfect'].setPlainText('새 문구');panel.category.setCurrentIndex(2);panel.category.setCurrentIndex(0)
    assert panel.editors['perfect'].toPlainText()=='새 문구'
    assert [w.settings_tabs.tabText(i) for i in range(3)]==['성적표 설정','자료 관리','사용 안내']
    bar=ManduProgressBar();bar.show();app.processEvents();assert bar.timer.isActive() and not bar.sprite.isNull()
    bar.hide();assert not bar.timer.isActive();bar.deleteLater()

def test_key_change_impact_and_points_only(fixture,tmp_path):
    from localomr.exam_changes import preview_change
    app,db,w,e,student,c,warnings=fixture
    sid=add_scan(db,e,student,tmp_path);db.confirm(sid);config=db.exam(e)['config']
    lines,ids=preview_change(db,e,[4,2,3],[2,3,4],config);assert ids==[sid] and '1번' in lines[0]
    lines,ids=preview_change(db,e,[1,2,3],[3,3,4],config);assert ids==[] and '배점' in lines[0]

def test_workbook_roundtrip_all_fields_and_partial_electives(fixture,tmp_path):
    from localomr.examdata import NUMERIC,effective
    from localomr.question_export import export_question_book
    from localomr.question_import import read_question_book,combine_sheets,HEADERS
    from openpyxl import load_workbook
    app,db,w,e,student,c,warnings=fixture
    key=[0 if q in NUMERIC else q%5+1 for q in range(1,31)]
    cfg={'form':'math30','primary':'1','topics':['단원']*30,'areas':['수학Ⅰ']*22+['확률과 통계']*8,'skills':['내용']*30,'difficulties':['중']*30,'conditions':['and']*30,'variants':{'2':{'key':key[22:],'points':[3.25]*8,'topics':['미분']*8,'areas':['미적분']*8,'skills':['도함수']*8,'difficulties':['상']*8,'conditions':['and']*8}}}
    eid=db.save_exam('가상 수학 1회','가상수학','수학','1','2026-10-20',key,[3.25]*30,config=cfg)
    exam=db.exam(eid);path=tmp_path/'roundtrip.xlsx';before=revision(db);export_question_book(path,exam)
    book=load_workbook(path);assert book.sheetnames==['확률과 통계','미적분'];assert [c.value for c in book.active[1]]==HEADERS;book.close()
    sheets=read_question_book(path);assert all(not s.errors for s in sheets)
    combined=combine_sheets(sheets)
    for choice,rows in combined.items():
        original=effective(exam,choice);assert len(rows)==30
        for i,row in enumerate(rows):
            assert row['q']==i+1 and row['key']==original['key'][i] and float(row['points'])==original['points'][i]/100
            for field,array in [('topic','topics'),('area','areas'),('skill','skills'),('difficulty','difficulties'),('kind','kinds'),('condition','conditions')]:assert row[field]==original[array][i]
    assert revision(db)==before

def test_exit_cancelled_backup_stays_open_and_uses_mascot(fixture,monkeypatch):
    app,db,w,e,student,c,warnings=fixture
    db.academy('추가 변경');seen=[]
    monkeypatch.setattr('localomr.popups.ask',lambda *a:seen.append(a) or QMessageBox.StandardButton.Yes)
    monkeypatch.setattr(w,'backup',lambda **kw:False)
    event=QCloseEvent();w.closeEvent(event)
    assert not event.isAccepted() and seen[0][3]=='exit_mandu.png' and '아직 백업하지 않은 변경사항이 있어요.' in seen[0][2]
    assert db.students()

def test_actual_popup_close_does_not_mean_skip(fixture):
    from PySide6.QtCore import QTimer
    from localomr.popups import ask,IllustratedMessage
    app,db,w,e,student,c,warnings=fixture
    seen=[]
    def close():
        popup=app.activeModalWidget();seen.append(isinstance(popup,IllustratedMessage) and not popup.iconPixmap().isNull());popup.close()
    QTimer.singleShot(10,close)
    choice=ask(w,'작업 시작 · 백업 확인','지금 당장 백업파일을 가져와주세요!!','import_mandu.png',QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No)
    assert choice==QMessageBox.StandardButton.Cancel and seen==[True]

def test_restore_preview_cancel_and_success_reset_exam_selection(fixture,tmp_path,monkeypatch):
    app,db,w,e,student,c,warnings=fixture
    path=db.backup(tmp_path/'input.omrbackup');db.save_student_contact('추가','5555',class_id=c)
    w.score_exam.setCurrentIndex(w.score_exam.findData(e));before=revision(db);messages=[]
    monkeypatch.setattr(QFileDialog,'getOpenFileName',lambda *a:(str(path),''))
    monkeypatch.setattr(QMessageBox,'exec',lambda box:messages.append(box.text()) or QMessageBox.StandardButton.Cancel)
    assert w.restore() is False and revision(db)==before and not w.startup_backup_ready
    assert '-1명' in messages[0] and '백업 저장 시각' in messages[0]
    monkeypatch.setattr(QMessageBox,'exec',lambda box:QMessageBox.StandardButton.Yes)
    assert w.restore() is True and w.startup_backup_ready and w.score_exam.currentData() is None
    assert len(db.students())==2 and not w.handoff.needs_backup()

def test_duplicate_short_phone_and_explicit_separate_registration(fixture,monkeypatch):
    from localomr.ui import StudentDialog
    app,db,w,e,student,c,warnings=fixture
    dialog=StudentDialog(db,parent=w);dialog.name.setText('가상학생');dialog.phone.setText('5678');seen=[]
    def choose(box,text):
        seen.append(box.text());next(b for b in box.buttons() if b.text()==text).click();return 0
    monkeypatch.setattr(QMessageBox,'exec',lambda box:choose(box,'기존 학생에게 반 추가'))
    dialog.save();assert len(db.students())==2 and '가상학원 / 가상반' in seen[0]
    monkeypatch.setattr(QMessageBox,'exec',lambda box:choose(box,'별도 등록'))
    dialog.save();assert len(db.students())==3 and dialog.student_id!=student

def test_result_checkbox_mouse_click_uses_visible_rows(fixture,tmp_path):
    from PySide6.QtTest import QTest
    app,db,w,e,student,c,warnings=fixture
    sid=add_scan(db,e,student,tmp_path);db.confirm(sid);w.startup_backup_ready=True;w.refresh();w.nav.setCurrentRow(5);w.show();app.processEvents()
    grid=w.result_table;cell=grid.visualItemRect(grid.item(0,0))
    QTest.mouseClick(grid.viewport(),Qt.MouseButton.LeftButton,pos=cell.center());app.processEvents()
    assert grid.checked_ids()==[sid] and w.selection_count.text()=='선택 1건'
    QTest.mouseClick(grid.viewport(),Qt.MouseButton.LeftButton,pos=cell.center());app.processEvents()
    assert grid.checked_ids()==[]
