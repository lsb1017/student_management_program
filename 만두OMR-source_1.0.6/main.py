from pathlib import Path
import argparse
import os
import sys
import traceback


def main():
    parser=argparse.ArgumentParser(description='만두OMR - offline Windows OMR grading')
    parser.add_argument('--data-dir',type=Path)
    parser.add_argument('--smoke-test',type=Path,help='Write a packaged runtime smoke-test result and exit')
    parser.add_argument('--math-smoke-input',type=Path,help='Optional synthetic math sheet for packaged integration test')
    parser.add_argument('--excel-smoke-input',type=Path,help='Optional question workbook for packaged import test')
    args=parser.parse_args()
    if args.smoke_test:
        import json, uuid
        import cv2, numpy
        from localomr.reports import sheet_pdf
        from localomr.omr import load_pages,recognize
        args.smoke_test.parent.mkdir(parents=True,exist_ok=True)
        form=args.smoke_test.with_suffix('.pdf')
        sheet_pdf(form,5,'실행파일 검증','000123',[1,2,3,4,5])
        _,img=next(load_pages(form));r=recognize(img,5)
        result={'ok':not r.error and [x['answer'] for x in r.answers]==[1,2,3,4,5] and r.code=='000123','answers':[x['answer'] for x in r.answers],'code':r.code,'error':r.error,'frozen':bool(getattr(sys,'frozen',False))}
        # Exercise the packaged database, PDF writer and UI without touching operational data.
        os.environ['QT_QPA_PLATFORM']='offscreen'
        from PySide6.QtWidgets import QApplication
        from PySide6.QtCore import QDate
        from localomr.store import Store
        from localomr.ui import MainWindow,STYLE,ReviewDialog
        from localomr.importer import import_files
        from localomr.reports import report_pdf
        smoke_dir=args.smoke_test.parent/('smoke_'+uuid.uuid4().hex[:8])
        smoke_dir.mkdir()
        db=Store(smoke_dir/'test.sqlite3')
        aid=db.academy('검증학원');cid=db.classroom(aid,'검증반');student=db.save_student('123','검증학생',cid)
        eid=db.save_exam('검증시험','검증','수학','1','2026-10-18',[1,2,3,4,5],[20]*5)
        imported=import_files(db.path,eid,[form])
        sid=db.scans(eid)[0]['id'];row=db.scan(sid)
        db.save_review(sid,student,row['answers'],[False]*5,True,'실행 검증');db.confirm(sid)
        assert not any(db.scan(sid)['reviewed']);result['unchecked_confirmation']=True
        report_pdf(smoke_dir/'report.pdf',db,[sid]);db.backup(smoke_dir/'backup.omrbackup')
        result['database_score']=db.results(eid)[0]['score']/100
        app=QApplication([]);app.setStyle('Fusion');app.setStyleSheet(STYLE)
        window=MainWindow(db);window.startup_backup_ready=True;window.score_exam.setCurrentIndex(window.score_exam.findData(eid));window.show();app.processEvents()
        try:window.selected_scoring_context()
        except ValueError:pass
        else:raise AssertionError('Import must require an academy and class')
        window.score_academy.setCurrentIndex(window.score_academy.findData(aid));window.score_class.setCurrentIndex(window.score_class.findData(cid))
        assert window.selected_scoring_context()==(aid,cid)
        result['required_import_scope']=True
        assert window.score_sat_on.date()==QDate.currentDate()
        window.grab().save(str(args.smoke_test.with_suffix('.png')))
        dialog=ReviewDialog(db,sid,window);dialog.resize(1200,800);dialog.show();app.processEvents()
        assert not hasattr(window.review_pane,'checked') and not hasattr(dialog,'review_widgets')
        assert dialog.grid.columnCount()==6 and dialog.filter.findText('미검수만')==-1
        result['question_review_controls_removed']=True
        dialog.close()
        if args.math_smoke_input:
            from localomr.mathomr import recognize_math
            from localomr.examdata import NUMERIC
            from localomr.examui import ExamEditor
            from PIL import Image
            expected=[(q%5)+1 if q not in NUMERIC else {16:1,17:0,18:12,19:100,20:288,21:47,22:17,29:844,30:622}[q] for q in range(1,31)]
            math_read=recognize_math(Image.open(args.math_smoke_input))
            assert [a['answer'] for a in math_read.answers]==expected and math_read.code=='042411'
            ms=db.save_student('42411','수학검증',cid)
            me=db.save_exam('수학검증','수학검증','수학','1','2026-10-18',expected,[3]*30,config={'form':'math30','primary':'1','cuts':{'1':[83,74,66,56,36,25,15,5]}})
            imp=import_files(db.path,me,[args.math_smoke_input],sat_on='2026-10-15');assert imp['added']==1 and not imp['failed']
            mid=db.scans(me)[0]['id'];mr=db.scan(mid);db.save_review(mid,ms,mr['answers'],[True]*30,True,elective='1');db.confirm(mid)
            report_pdf(smoke_dir/'math_report.pdf',db,[mid])
            import pypdfium2 as pdfium
            with pdfium.PdfDocument(str(smoke_dir/'math_report.pdf')) as doc:
                assert len(doc)==2
                page=doc[0];tp=page.get_textpage();report_text=tp.get_text_range();tp.close()
                assert any(obj.type==pdfium.raw.FPDF_PAGEOBJ_IMAGE for obj in page.get_objects())
                page.close()
                for expected_text in ('양지용 수학연구실','학생정보','단원별 성취도 분석','학원명','공통 / 선택과목 점수','확률과 통계 24점','만점 달성!','완벽하게 해냈어요!'):assert expected_text in report_text
                assert '응시 2026-10-15' not in report_text
                for expected_text in ('클리닉 우선순위 분석','이번 시험은 모든 문항을 맞혔어요!','오답이 없습니다.'):
                    assert expected_text in report_text
                page=doc[1];tp=page.get_textpage();report_text=tp.get_text_range();tp.close();page.close()
                for expected_text in ('문항별 채점표','오답률 상위 분석','전체 오답률','성적 추이 분석','응시일자'):assert expected_text in report_text
                assert report_text.count('2026-10-15')==2
                assert '시험명' in report_text and report_text.count('수학검증')==1
            assert b'NanumDaeGwangYuRi' in (smoke_dir/'math_report.pdf').read_bytes()
            result['report_design_assets']=True
            result['report_insights']=True
            from localomr.store import pack
            from localomr.historyui import HistoryDialog
            with db.transaction():
                he=db.db.execute("INSERT INTO historical_exams(title,day,mean,maximum,source_name,source_hash) VALUES('이전 검증시험 1회','2026-09-01',72.94,10000,'synthetic.xlsx','synthetic-history')").lastrowid
                db.db.execute('INSERT INTO historical_results(exam_id,student_id,sat_on,score,grade,elective,snapshot,earned,ox,answers,source_row) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(he,ms,'2026-09-01',8400,'2','1',pack(db.student(ms)),pack([200]*22+[500]*8),pack(['O']*30),pack(expected),5))
            history_dialog=HistoryDialog(db,window);assert history_dialog.table.rowCount()==1;history_dialog.close()
            report_pdf(smoke_dir/'history_report.pdf',db,[mid])
            with pdfium.PdfDocument(str(smoke_dir/'history_report.pdf')) as doc:
                page=doc[1];tp=page.get_textpage();history_text=tp.get_text_range();tp.close();page.close()
                assert '이전 검증시험 1회' in history_text and '72.9' in history_text
            result['historical_scores_and_trends']=True
            from PIL import ImageDraw
            from localomr.mathomr import MARKERS
            from localomr.omr import ReadResult,png_bytes,pixel_hash
            from localomr.importer import retry_failed
            gray=Image.open(args.math_smoke_input).convert('RGB');draw=ImageDraw.Draw(gray)
            for i,(x,y) in enumerate(MARKERS):draw.rectangle((x-8,y-(16 if i<2 else 8),x+8,y+(16 if i<2 else 8)),fill=(145,165,150))
            gray_path=smoke_dir/'gray_markers.png';gray.save(gray_path)
            failed=ReadResult(None,[{'answer':None,'flag':'failed','levels':[]} for _ in expected],None,error='previous marker threshold')
            retry_id,_=db.add_scan(me,db.source(gray_path),1,pixel_hash(gray),png_bytes(gray),failed,sat_on='2026-10-12')
            recovered=retry_failed(db.path,me,[retry_id],class_id=cid)
            retried=db.scan(retry_id)
            assert recovered['recovered']==1 and retried['answers']==expected and retried['status']=='pending' and retried['sat_on']=='2026-10-12'
            assert db.scan(mid)['status']=='confirmed'
            result['failed_scan_recovery']=True
            from localomr.mathomr import question_rect,name_centers,NAME_CONSONANTS,NAME_VOWELS,HANGUL_INITIALS,HANGUL_VOWELS,HANGUL_FINALS
            window.refresh();window.nav.setCurrentRow(3);window.score_exam.setCurrentIndex(window.score_exam.findData(me));app.processEvents()
            index=window.visible_scan_ids.index(retry_id);window.scan_table.setCurrentCell(index,11)
            assert window.review_pane.view.highlight.rect().getRect()==question_rect(7)
            window.scan_table.item(index,11).setText('1')
            assert db.scan(retry_id)['answers'][6]==1 and '×' in window.scan_table.item(index,11).text()
            assert db.scan(mid)['status']=='confirmed'
            result['inline_keyboard_review']=True
            from PySide6.QtWidgets import QLineEdit
            from PySide6.QtTest import QTest
            from PySide6.QtCore import Qt
            window.activateWindow();window.scan_table.setCurrentCell(index,2);window.scan_table.scrollToItem(window.scan_table.item(index,2));window.scan_table.setFocus();app.processEvents();QTest.keyClicks(window.scan_table,'2');app.processEvents()
            editor=app.focusWidget();assert isinstance(editor,QLineEdit)
            QTest.keyClick(editor,Qt.Key.Key_Right);app.processEvents();app.processEvents()
            assert window.scan_table.cellWidget(index,2) is None and db.scan(retry_id)['elective']=='2'
            window.scan_table.item(index,2).setText('1');app.processEvents()
            assert db.scan(retry_id)['elective']=='1'
            result['direct_elective_input']=True
            from localomr.importer import reread_selected
            assert reread_selected(db.path,me,[retry_id],academy_id=aid,class_id=cid)['recovered']==1
            assert db.scan(retry_id)['answers'][6]==1 and db.scan(mid)['status']=='confirmed'
            result['reread_preserves_edits']=True
            named=gray.copy();draw=ImageDraw.Draw(named)
            for s,ch in enumerate('김민수'):
                number=ord(ch)-0xAC00;parts=[HANGUL_INITIALS[number//588],HANGUL_VOWELS[number//28%21],HANGUL_FINALS[number%28]]
                for p,ch in enumerate(parts):
                    if ch==' ':continue
                    labels=NAME_VOWELS if p==1 else NAME_CONSONANTS;x,y=name_centers(s,p)[labels.index(ch)];draw.ellipse((x-5,y-11,x+5,y+11),fill='black')
            named_student=db.save_student_contact('김민수','4241',class_id=cid)
            other_class=db.classroom(db.academy('다른 검증학원'),'다른반');db.save_student_contact('김민수','4241',class_id=other_class)
            name_path=smoke_dir/'named.png';named.save(name_path)
            assert import_files(db.path,me,[name_path],academy_id=aid,class_id=cid)['added']==1
            named_id=db.scans(me)[-1]['id'];assert db.scan(named_id)['student_id']==named_student and not db.scan(named_id)['identity_ok']
            result['phone_name_scoped_identity']=True
            assert db.delete_scans(me,[named_id])==1 and named_id not in [s['id'] for s in db.scans(me)]
            assert db.restore_deleted_scans(me)==1 and db.scan(named_id)['student_id']==named_student
            result['scan_delete_undo']=True
            md=ReviewDialog(db,mid,window);md.resize(1400,900);md.show();app.processEvents();md.grab().save(str(smoke_dir/'math_review.png'));md.close()
            ed=ExamEditor(db,me,window);ed.show();app.processEvents();ed.close()
            result['math']={'answers':30,'score':db.results(me)[0]['score']/100,'grade':db.results(me)[0]['grade'],'estimated_percentile':db.results(me)[0]['estimated_percentile'],'report_pages':2}
            from localomr.batchui import BatchOutputDialog
            from localomr.ranking import ranking_xlsx
            import openpyxl
            batch=BatchOutputDialog(db,me,window);batch.show();app.processEvents();assert len(batch.rows)==1;batch.grab().save(str(smoke_dir/'batch_preview.png'));batch.close()
            ranking_xlsx(smoke_dir/'ranks.xlsx',db,me,[cid])
            workbook=openpyxl.load_workbook(smoke_dir/'ranks.xlsx');assert workbook.active['K9'].value==90;workbook.close();result['batch_and_rank_xlsx']=True
        if args.excel_smoke_input:
            from localomr.question_import import read_question_book,combine_sheets
            from localomr.questionui import QuestionImportDialog
            from localomr.examui import ExamEditor
            sheets=read_question_book(args.excel_smoke_input)
            assert len(sheets)==8 and all(not s.errors and len(s.rows)==30 for s in sheets)
            selected=[s for s in sheets if s.group=='시즌2-2'];assert len(combine_sheets(selected))==3
            ed=ExamEditor(db);ed.apply_import(selected,args.excel_smoke_input.name);ed.series.setText('엑셀 검증');ed.round.setText('1');ed.save();assert ed.eid
            preview=QuestionImportDialog(sheets,window);preview.show();app.processEvents();preview.grab().save(str(smoke_dir/'excel_preview.png'));preview.close()
            result['excel']={'sheets':8,'questions':240,'imported_electives':3,'common_conflict_detected':False}
            try:combine_sheets([s for s in sheets if s.group=='시즌2-1'])
            except ValueError:result['excel']['common_conflict_detected']=True
            assert result['excel']['common_conflict_detected']
        from localomr.studentdata import template_path,read_students
        assert template_path().is_file()
        ns=db.save_student_contact('수기검증','010-5000-6789',school='예제고',parent_phone='01050001234',class_id=cid)
        added,skipped=db.import_students([{'name':'엑셀학생','phone':'6789'},{'name':'엑셀학생','phone':'6789'}])
        assert len(added)==1 and skipped==1 and db.student(added[0])['code']=='067892'
        # Verify and acknowledge the new navigation guard in the packaged smoke run.
        from PySide6.QtCore import QTimer
        from PySide6.QtWidgets import QMessageBox
        from localomr.popups import PendingReviewDialog
        def acknowledge_pending_navigation():
            dialog=app.activeModalWidget()
            if isinstance(dialog,PendingReviewDialog):
                dialog.grab().save(str(smoke_dir/'pending_navigation.png'));result['pending_navigation_dialog']=True;dialog.done(int(QMessageBox.StandardButton.Yes))
        window.refresh();window.manual.open_student(eid,ns);QTimer.singleShot(0,acknowledge_pending_navigation);window.nav.setCurrentRow(4)
        assert not any(window.manual.state['entered'])
        assert window.manual.sat_on.date()==QDate.currentDate()
        window.manual.sat_on.setDate(QDate(2026,10,14))
        window.manual.toggle(0);assert window.manual.state['answers'][0]==1
        window.manual.full_marks();state=window.manual.state
        mid=db.save_manual(eid,ns,None,state['answers'],state['entered'],state['revision'],sat_on=state['sat_on'])
        window.manual.discard_all();window.refresh();window.manual.open_student(eid,ns);app.processEvents()
        assert db.results(eid,ns)[0]['score']==10000
        assert db.results(eid,ns)[0]['day']=='2026-10-14' and window.manual.sat_on.date()==QDate(2026,10,14)
        result['attendance_dates']=True
        window.grab().save(str(smoke_dir/'manual_preview.png'))
        db.backup(smoke_dir/'with_manual.omrbackup')
        result['manual_and_students']={'blank_unattempted':True,'ox_autofill':True,'manual_score':100,'phone_code':db.student(ns)['code'],'batch_added':1,'batch_duplicate_skipped':1,'template_embedded':True}
        assert window.report_brand.text()=='양지용 수학연구실'
        result['research_name_setting']=True
        from localomr.cheerui import CheerDialog
        from localomr import cheer
        from PySide6.QtCore import Qt
        cheer_dialog=CheerDialog(db,window);cheer_dialog.show();app.processEvents()
        assert all(not view.pixmap().isNull() for view in cheer_dialog.previews.values())
        cheer_dialog.editors['grade1'].setPlainText('1등급 축하해요!')
        cheer_dialog.save();db.backup(smoke_dir/'cheer.omrbackup')
        assert cheer.settings(db)['grade1']['text']=='1등급 축하해요!'
        window.nav.setCurrentRow(5);window.result_table.set_checks(True)
        assert set(window.result_table.checked_ids())=={r['id'] for r in window.filtered_results()}
        window.result_table.sortItems(7,Qt.SortOrder.DescendingOrder)
        assert window.student_table.horizontalHeaderItem(0).text()=='학원'
        db.archive_exams([eid]);assert not db.results(eid)
        db.restore_exams([eid]);assert len(db.results(eid))==2
        result['custom_cheer_font_images']=True
        result['result_checks_sorting']=True
        result['reversible_exam_delete']=True
        if args.math_smoke_input:
            from PySide6.QtWidgets import QMessageBox
            db.save_review(named_id,named_student,expected,[False]*30,True,elective='1')
            window.refresh();window.score_exam.setCurrentIndex(window.score_exam.findData(me))
            window.score_academy.setCurrentIndex(window.score_academy.findData(aid));window.score_class.setCurrentIndex(window.score_class.findData(cid))
            old_question,old_info=QMessageBox.question,QMessageBox.information
            QMessageBox.question=lambda *a:QMessageBox.StandardButton.Yes;QMessageBox.information=lambda *a:None
            def acknowledge_result():
                box=app.activeModalWidget()
                if isinstance(box,QMessageBox) and box.windowTitle()=='일괄 확정 결과':
                    assert box.property('mascotAsset')=='frustrated_mandu.png'
                    box.grab().save(str(smoke_dir/'bulk_partial_result.png'))
                    result['conditional_result_mascot']=True
                    box.accept()
            QTimer.singleShot(0,acknowledge_result)
            try:window.confirm_visible_scans()
            finally:QMessageBox.question, QMessageBox.information=old_question,old_info
            assert db.scan(named_id)['status']=='confirmed' and db.scan(retry_id)['status']=='pending'
            result['scoped_bulk_confirmation']=True
        from localomr.organizationui import OrganizationDialog
        from localomr.catalog import reference
        org=OrganizationDialog(db,window);org.show();app.processEvents();org.close()
        db.rename_organization('academy',aid,'변경 검증학원');db.rename_organization('class',cid,'변경 검증반')
        assert db.student(student)['class_id']==cid and db.scan(sid)['snapshot']['academy']=='변경 검증학원'
        assert reference('2027 수능 Fit × 양지용 6회차')['cuts']['2']==[87,82,65,53,41]
        result['organization_rename']=True;result['supplied_cutoff_catalog']=True
        from localomr import __version__
        from localomr.storageui import StorageDialog
        from localomr.makeupui import MakeupStudentDialog
        from localomr.preview_cache import preview_pdf,cache_files
        from localomr.export_names import export_name
        from localomr import storage
        assert not window.windowIcon().isNull()
        assert hasattr(window.review_pane,'next_student_button') and not hasattr(window.review_pane,'identity')
        result['confirmation_without_checkbox']=True;result['heart_icon']=True;result['next_student_button']=True
        assert not window.home_mascot.pixmap().isNull()
        assert all(card.parentWidget().height()==94 for card in window.home_cards)
        result['compact_home_with_mascot']=True
        cleanup=StorageDialog(db,window);cleanup.show();app.processEvents();cleanup.close()
        makeup=MakeupStudentDialog(db,window);makeup.search.setText('검증');makeup.show();app.processEvents();assert makeup.table.rowCount()>0;makeup.close()
        result['storage_and_makeup_dialogs']=True
        assert 'grade5plus' in cheer.settings(db) and cheer.image_bytes('grade5plus',cheer.settings(db)['grade5plus'])
        result['four_cheer_categories']=True
        rows=db.results(eid);assert export_name(rows).endswith('.pdf') and export_name(rows,'rank').endswith('.xlsx')
        result['export_filename_rules']=True
        cached=preview_pdf(db,sid);assert preview_pdf(db,sid)==cached and len(cache_files(db))==1
        result['preview_cache_reuse']=True
        with db.transaction():db.db.execute("UPDATE scans SET created_at='2020-01-01T00:00:00+00:00' WHERE id=?",(sid,))
        before=db.results();cleaned=storage.purge(db,[sid])
        assert before==db.results() and not db.scan(sid)['raw'] and not cleaned['warning']
        result['image_cleanup_preserves_grades']=True;result['app_version']=__version__
        from localomr.backupui import BackupNameDialog
        from localomr.popups import PendingReviewDialog
        from localomr.busy import ManduProgressBar
        from localomr.reports import report_asset
        from PySide6.QtGui import QPixmap
        assert __version__=='1.0.6'
        naming=BackupNameDialog(db,{},window);naming.academy.setCurrentIndex(naming.academy.findData(aid));naming.classroom.setCurrentIndex(naming.classroom.findData(cid))
        assert naming.names()==('변경 검증학원','변경 검증반')
        assert '_변경 검증학원_변경 검증반' in naming.filename_preview.text();naming.show();app.processEvents();naming.hide()
        for asset in ('smart_mandu.png','celebrate_mandu.png','app_icon_heart.png','busy_mandu_alpha.png'):assert not QPixmap(str(report_asset(asset))).isNull()
        from localomr.popup_style import ASSETS,decorate
        for asset in set(ASSETS.values())|{'celebrate_mandu.png','running_mandu.png','arrow_mandu.png'}:
            pix=QPixmap(str(report_asset(asset)));assert not pix.isNull() and pix.hasAlphaChannel(),asset
        probe=QMessageBox(window);probe.setWindowTitle('선택 시험 삭제');probe.setText('삭제할까요?');probe.setStandardButtons(QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No)
        decorate(probe);assert probe.property('mascotAsset')=='crazy_mandu.png' and probe.defaultButton()==probe.button(QMessageBox.StandardButton.No)
        result['popup_assets_and_safe_defaults']=True
        from localomr.popup_style import popup_mascot_pixmap
        assert popup_mascot_pixmap('thumb_mandu.png').height()<86
        assert not window.pending_mascot.icon().isNull()
        result['small_popup_mascots_and_pending_arrow']=True
        assert cleanup.days.minimum()==0
        cleanup.days.setValue(0);assert cleanup.days.value()==0
        cleanup.broaden();assert cleanup.days.value()==2
        bar=ManduProgressBar(window);assert bar.sprite.height()==60
        assert window.backup_button.parentWidget()==window.home_notice and window.restore_button.parentWidget()==window.home_notice
        result['version_1_assets_and_backup_dialog']=True;result['zero_day_cleanup_and_two_day_shortcut']=True
        result['large_import_limits']={'bytes':__import__('localomr.omr',fromlist=['MAX_FILE_BYTES']).MAX_FILE_BYTES,'pages':__import__('localomr.omr',fromlist=['MAX_PAGES']).MAX_PAGES}
        extra_class=db.classroom(aid,'여러반 검증')
        db.set_student_classes(student,[cid,extra_class]);assert db.in_class(student,extra_class)
        assert db.db.execute('PRAGMA user_version').fetchone()[0]==10
        assert window.result_choice.isHidden()
        from localomr.drive_backup import DriveClient,protect
        drive_probe=DriveClient(smoke_dir/'drive-probe',report_asset('google_client.json'))
        assert drive_probe.client()['client_id']
        drive_probe.save_token({'access_token':'synthetic','refresh_token':'synthetic'})
        assert b'synthetic' not in drive_probe.token_path.read_bytes()
        result['multiple_class_membership']=True;result['drive_client_and_encrypted_credentials']=True
        result['drive_live_account_verified']=False
        window.handoff_reminder.setChecked(False);window.close()
        from localomr.forms import math_form_path
        assert math_form_path().stat().st_size>100000
        result['default_math_form']=True
        result['ui']='home_and_review_rendered'
        result['ok']=result['ok'] and result['database_score']==100 and imported['added']==1
        args.smoke_test.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        return 0 if result['ok'] else 1
    from PySide6.QtCore import QLockFile
    from PySide6.QtWidgets import QApplication,QMessageBox
    from localomr.store import Store
    from localomr.ui import MainWindow,STYLE
    app=QApplication(sys.argv);app.setApplicationName('만두OMR');app.setStyle('Fusion');app.setStyleSheet(STYLE)
    from localomr.popup_style import install
    install()
    # Data next to a portable EXE; --data-dir supports an explicitly selected location.
    base=Path(sys.executable).parent if getattr(sys,'frozen',False) else Path(__file__).resolve().parent
    data_dir=(args.data_dir or base/'data').resolve()
    try:
        data_dir.mkdir(parents=True,exist_ok=True)
        lock=QLockFile(str(data_dir/'session.lock'));lock.setStaleLockTime(0)
        if not lock.tryLock(100):
            QMessageBox.warning(None,'이미 실행 중','이 데이터 폴더를 사용하는 만두OMR가 이미 실행 중입니다.');return 1
        store=Store(data_dir/'localomr.sqlite3')
        window=MainWindow(store);window.show()
        def report_error(kind,value,tb):
            (data_dir/'last_error.log').write_text(''.join(traceback.format_exception(kind,value,tb)),encoding='utf-8')
            QMessageBox.warning(window,'처리 오류',str(value))
        sys.excepthook=report_error
        return app.exec()
    except Exception as exc:
        QMessageBox.critical(None,'시작할 수 없습니다',f'{exc}\n쓰기 가능한 폴더에 압축을 풀어 실행하세요.');return 1


if __name__=='__main__':
    try:
        sys.exit(main())
    except Exception:
        if '--smoke-test' in sys.argv:
            failure=Path(sys.argv[sys.argv.index('--smoke-test')+1])
            failure.parent.mkdir(parents=True,exist_ok=True)
            failure.write_text(traceback.format_exc(),encoding='utf-8')
            sys.exit(1)
        raise
